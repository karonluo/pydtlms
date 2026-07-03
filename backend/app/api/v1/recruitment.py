from datetime import datetime
from io import BytesIO
from pathlib import Path
from urllib.parse import quote
from uuid import uuid4
from typing import Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, Path as PathParam, Query, Response, UploadFile, status
from fastapi.responses import StreamingResponse

from app.core.rbac import require_permissions
from app.schemas.auth import Principal
from app.schemas.recruitment import (
    AdvisorScreeningScoreUpdateRequest,
    AdvisorScreeningBatchSubmitRequest,
    AdvisorScreeningBatchSubmitResponse,
    HackathonScoreImportResult,
    CampOfferImportResult,
    CampOfferListResponse,
    CampOfferStats,
    CampOfferNotificationSendRequest,
    CampOfferNotificationSendResponse,
    CampOfferRecord,
    CampOfferUpsert,
    OfferTemplateListResponse,
    OfferTemplateRecord,
    InitialScreeningConfirmationApplicationListResponse,
    InitialScreeningConfirmationApplicationRecord,
    InitialScreeningConfirmationRequest,
    RecruitApplicationListResponse,
    RecruitApplicationImportResult,
    RecruitApplicationRecord,
    RecruitPortalApplicationDetail,
    RecruitApplicationUpsert,
    RecruitPlanListResponse,
    RecruitPlanRecord,
    RecruitPlanUpsert,
    RecruitmentOptionsResponse,
    RecruitStats,
    RecruitWorkbench,
)
from app.schemas.student import RegisteredPortalStudentExportJobCreateResponse, RegisteredPortalStudentExportRequest
from app.services.dashboard_service import (
    confirm_initial_screening,
    create_camp_offer,
    create_recruitment_application,
    create_recruitment_plan,
    create_registered_portal_student_export_job,
    delete_camp_offer,
    delete_recruitment_plan,
    delete_recruitment_application,
    export_camp_offers,
    export_recruitment_application_blank_template,
    export_recruitment_applications,
    get_recruitment_application_detail,
    get_recruitment_application_list,
    get_recruitment_options,
    get_recruitment_portal_application_detail,
    get_recruitment_plan_list,
    get_recruitment_stats,
    get_recruitment_workbench,
    get_camp_offer_detail,
    get_camp_offer_list,
    get_camp_offer_stats,
    import_recruitment_applications,
    rescore_advisor_screening_submitted_application,
    submit_advisor_screening_batch,
    update_camp_offer,
    update_recruitment_application,
    update_recruitment_plan,
    update_advisor_screening_score,
    set_camp_offer_accepted_status,
)
from app.services.advisor_screening_pending_service import list_advisor_screening_pending_applications
from app.services.advisor_screening_submitted_service import (
    count_advisor_screening_submitted_applications,
    list_advisor_screening_submitted_applications,
    query_store as advisor_screening_submitted_query_store,
)
from app.services.initial_screening_confirmation_service import list_initial_screening_confirmation_applications
from app.services.camp_offer_confirmation_service import submit_camp_offer_confirmation
from app.services.camp_offer_import_service import (
    import_camp_offers_from_excel,
    import_hackathon_scores_from_excel,
)
from app.services.camp_offer_notification_service import (
    OFFER_TEMPLATE_UPLOAD_DIR as _OFFER_TEMPLATE_UPLOAD_DIR,
    send_camp_offer_notifications,
    _resolve_template_path as resolve_offer_template_path,
    _builtin_template_path as resolve_builtin_offer_template_path,
)
from app.services.offer_template_renderer import (
    render_with_sample_placeholders,
    render_markdown_to_html,
)
from app.services.recruitment_excel_service import parse_recruitment_template


router = APIRouter(prefix="/recruitment", tags=["recruitment"])
PROJECT_ROOT = Path(__file__).resolve().parents[4]
BROCHURE_UPLOAD_DIR = PROJECT_ROOT / "frontend" / "public" / "portal-brochures" / "uploads"
OFFER_TEMPLATE_UPLOAD_DIR = _OFFER_TEMPLATE_UPLOAD_DIR
OFFER_TEMPLATE_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
OFFER_TEMPLATE_MAX_BYTES = 1 * 1024 * 1024
OFFER_TEMPLATE_ALLOWED_SUFFIXES = {".md", ".markdown"}
BUILTIN_OFFER_TEMPLATE_LABEL = {
    "first": "系统内置 - 第一志愿（offer.md）",
    "second": "系统内置 - 第二志愿（offer2.md）",
}
BUILTIN_OFFER_TEMPLATE_FILENAME = {
    "first": "offer.md",
    "second": "offer2.md",
}


@router.get("/workbench", response_model=RecruitWorkbench)
def recruitment_workbench(principal: Principal = Depends(require_permissions("recruitment:read"))) -> RecruitWorkbench:
    return get_recruitment_workbench()


@router.get("/stats", response_model=RecruitStats)
def recruitment_stats(principal: Principal = Depends(require_permissions("recruitment:read"))) -> RecruitStats:
    return get_recruitment_stats()


@router.get("/options", response_model=RecruitmentOptionsResponse)
def recruitment_options(principal: Principal = Depends(require_permissions("recruitment:read"))) -> RecruitmentOptionsResponse:
    return get_recruitment_options()


@router.get("/plans", response_model=RecruitPlanListResponse)
def recruitment_plans(
    keyword: str | None = Query(default=None),
    semester: str | None = Query(default=None),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=10, ge=1, le=1000),
    principal: Principal = Depends(require_permissions("recruitment:read")),
) -> RecruitPlanListResponse:
    return get_recruitment_plan_list(keyword=keyword, semester=semester, page=page, page_size=page_size)


@router.post("/plans", response_model=RecruitPlanRecord, status_code=status.HTTP_201_CREATED)
def create_recruitment_plan_record(payload: RecruitPlanUpsert, principal: Principal = Depends(require_permissions("recruitment:write"))) -> RecruitPlanRecord:
    return create_recruitment_plan(payload)


@router.put("/plans/{plan_id}", response_model=RecruitPlanRecord)
def update_recruitment_plan_record(plan_id: int, payload: RecruitPlanUpsert, principal: Principal = Depends(require_permissions("recruitment:write"))) -> RecruitPlanRecord:
    try:
        return update_recruitment_plan(plan_id, payload)
    except KeyError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Recruitment plan not found") from exc


@router.delete("/plans/{plan_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_recruitment_plan_record(plan_id: int, principal: Principal = Depends(require_permissions("recruitment:write"))) -> None:
    try:
        delete_recruitment_plan(plan_id)
    except KeyError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Recruitment plan not found") from exc
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc


@router.post("/plans/brochure-upload")
async def upload_recruitment_brochure_image(
    file: UploadFile = File(...),
    principal: Principal = Depends(require_permissions("recruitment:write")),
) -> dict[str, str]:
    content_type = str(file.content_type or "")
    if not content_type.startswith("image/"):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="仅支持上传图片文件")

    suffix = Path(file.filename or "brochure.png").suffix.lower() or ".png"
    if suffix not in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".svg"}:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="图片格式不受支持")

    BROCHURE_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    filename = f"brochure-{uuid4().hex}{suffix}"
    target = BROCHURE_UPLOAD_DIR / filename
    target.write_bytes(await file.read())
    return {"url": f"/portal-brochures/uploads/{filename}"}


@router.get("/applications", response_model=RecruitApplicationListResponse)
def recruitment_applications(
    keyword: str | None = Query(default=None),
    plan_id: int | None = Query(default=None),
    status_filter: str | None = Query(default=None, alias="status"),
    portal_student_only: bool = Query(default=False),
    advisor_names_filter: str | None = Query(default=None, alias="advisor_names"),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=10, ge=1, le=1000),
    principal: Principal = Depends(require_permissions("recruitment:read")),
) -> RecruitApplicationListResponse:
    advisor_names = [segment.strip() for segment in str(advisor_names_filter or "").split(",") if segment.strip()]
    return get_recruitment_application_list(
        keyword=keyword,
        plan_id=plan_id,
        status=status_filter,
        portal_student_only=portal_student_only,
        advisor_names=advisor_names or None,
        principal=principal,
        page=page,
        page_size=page_size,
    )


@router.get("/applications/initial-screening-confirmation", response_model=InitialScreeningConfirmationApplicationListResponse)
def recruitment_initial_screening_confirmation_applications(
    keyword: str | None = Query(default=None),
    plan_id: int = Query(..., ge=1),
    advisor_names_filter: str | None = Query(default=None, alias="advisor_names"),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=10, ge=1, le=1000),
    principal: Principal = Depends(require_permissions("recruitment:read")),
) -> InitialScreeningConfirmationApplicationListResponse:
    advisor_names = [segment.strip() for segment in str(advisor_names_filter or "").split(",") if segment.strip()]
    return list_initial_screening_confirmation_applications(
        keyword=keyword,
        plan_id=plan_id,
        advisor_names=advisor_names or None,
        page=page,
        page_size=page_size,
    )


@router.get("/applications/advisor-screening-submitted", response_model=InitialScreeningConfirmationApplicationListResponse)
def recruitment_advisor_screening_submitted_applications(
    keyword: str | None = Query(default=None),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=10, ge=1, le=1000),
    principal: Principal = Depends(require_permissions("recruitment_advisor_screening:read")),
) -> InitialScreeningConfirmationApplicationListResponse:
    advisor_name = principal.full_name.strip() or None
    advisor_user_id = advisor_screening_submitted_query_store._advisor_user_id_by_username(principal.username)
    return list_advisor_screening_submitted_applications(
        keyword=keyword,
        advisor_name=advisor_name,
        advisor_user_id=advisor_user_id,
        page=page,
        page_size=page_size,
    )


@router.get("/applications/advisor-screening-submitted/count")
def recruitment_advisor_screening_submitted_count(
    keyword: str | None = Query(default=None),
    principal: Principal = Depends(require_permissions("recruitment_advisor_screening:read")),
) -> dict[str, int]:
    advisor_name = principal.full_name.strip() or None
    advisor_user_id = advisor_screening_submitted_query_store._advisor_user_id_by_username(principal.username)
    return {
        "total": count_advisor_screening_submitted_applications(
            keyword=keyword,
            advisor_name=advisor_name,
            advisor_user_id=advisor_user_id,
        )
    }


@router.get("/applications/advisor-screening-pending")
def recruitment_advisor_screening_pending_applications(
    keyword: str | None = Query(default=None),
    principal: Principal = Depends(require_permissions("recruitment_advisor_screening:read")),
) -> list[dict[str, Any]]:
    advisor_name = principal.full_name.strip() or None
    advisor_user_id = None
    return list_advisor_screening_pending_applications(
        keyword=keyword,
        advisor_username=principal.username,
        advisor_name=advisor_name,
        advisor_user_id=advisor_user_id,
    )


@router.get("/applications/advisor-screening-pending/count")
def recruitment_advisor_screening_pending_count(
    keyword: str | None = Query(default=None),
    principal: Principal = Depends(require_permissions("recruitment_advisor_screening:read")),
) -> dict[str, int]:
    advisor_name = principal.full_name.strip() or None
    advisor_user_id = None
    return {
        "total": len(
            list_advisor_screening_pending_applications(
                keyword=keyword,
                advisor_username=principal.username,
                advisor_name=advisor_name,
                advisor_user_id=advisor_user_id,
            )
        )
    }


@router.post("/applications", response_model=RecruitApplicationRecord, status_code=status.HTTP_201_CREATED)
def create_recruitment_application_record(payload: RecruitApplicationUpsert, principal: Principal = Depends(require_permissions("recruitment:write"))) -> RecruitApplicationRecord:
    return create_recruitment_application(payload, principal=principal)


@router.put("/applications/{application_id}", response_model=RecruitApplicationRecord)
def update_recruitment_application_record(application_id: int, payload: RecruitApplicationUpsert, principal: Principal = Depends(require_permissions("recruitment:write"))) -> RecruitApplicationRecord:
    try:
        return update_recruitment_application(application_id, payload)
    except KeyError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Recruitment application not found") from exc
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc


@router.post("/applications/advisor-screening:submit", response_model=AdvisorScreeningBatchSubmitResponse)
def submit_advisor_screening_batch_record(
    payload: AdvisorScreeningBatchSubmitRequest,
    principal: Principal = Depends(require_permissions("recruitment:write")),
) -> AdvisorScreeningBatchSubmitResponse:
    try:
        return submit_advisor_screening_batch(payload, principal=principal)
    except KeyError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Recruitment application not found") from exc
    except PermissionError as exc:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc


@router.post("/applications/advisor-screening-score", response_model=RecruitApplicationRecord)
def update_advisor_screening_score_record(
    payload: AdvisorScreeningScoreUpdateRequest,
    principal: Principal = Depends(require_permissions("recruitment:write")),
) -> RecruitApplicationRecord:
    try:
        return update_advisor_screening_score(payload, principal=principal)
    except KeyError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Recruitment application not found") from exc
    except PermissionError as exc:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc


@router.post("/applications/{application_id}/initial-screening-confirmation", response_model=RecruitApplicationRecord)
def confirm_initial_screening_record(
    application_id: int,
    payload: InitialScreeningConfirmationRequest,
    principal: Principal = Depends(require_permissions("recruitment:write")),
) -> RecruitApplicationRecord:
    try:
        return confirm_initial_screening(application_id, payload, principal=principal)
    except KeyError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Recruitment application not found") from exc
    except PermissionError as exc:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc


@router.post("/applications/advisor-screening-submitted/{application_id}/rescore", response_model=RecruitApplicationRecord)
def rescore_advisor_screening_submitted_record(
    application_id: int,
    principal: Principal = Depends(require_permissions("recruitment:write")),
) -> RecruitApplicationRecord:
    try:
        return rescore_advisor_screening_submitted_application(application_id, principal=principal)
    except KeyError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Recruitment application not found") from exc
    except PermissionError as exc:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
    except RuntimeError as exc:
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail=str(exc)) from exc


@router.get("/applications/export")
def export_recruitment_application_records(
    keyword: str | None = Query(default=None),
    plan_id: int | None = Query(default=None),
    status_filter: str | None = Query(default=None, alias="status"),
    portal_student_only: bool = Query(default=False),
    advisor_names_filter: str | None = Query(default=None, alias="advisor_names"),
    principal: Principal = Depends(require_permissions("recruitment:read")),
) -> StreamingResponse:
    advisor_names = [segment.strip() for segment in str(advisor_names_filter or "").split(",") if segment.strip()]
    content = export_recruitment_applications(
        keyword=keyword,
        plan_id=plan_id,
        status=status_filter,
        portal_student_only=portal_student_only,
        advisor_names=advisor_names or None,
        principal=principal,
    )
    filename = f"资料审核名单_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    encoded_filename = quote(filename)
    return StreamingResponse(
        iter([content]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )


@router.get("/applications/template")
def download_recruitment_application_template(
    principal: Principal = Depends(require_permissions("recruitment:read")),
) -> StreamingResponse:
    content = export_recruitment_application_blank_template()
    filename = "资料审核名单模板.xlsx"
    encoded_filename = quote(filename)
    return StreamingResponse(
        iter([content]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )


@router.get("/applications/{application_id}", response_model=RecruitApplicationRecord)
def recruitment_application_detail(application_id: int, principal: Principal = Depends(require_permissions("recruitment:read"))) -> RecruitApplicationRecord:
    try:
        return get_recruitment_application_detail(application_id)
    except KeyError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Recruitment application not found") from exc


@router.get("/applications/{application_id}/portal-detail", response_model=RecruitPortalApplicationDetail)
def recruitment_portal_application_detail(application_id: int, principal: Principal = Depends(require_permissions("recruitment:read"))) -> RecruitPortalApplicationDetail:
    try:
        return get_recruitment_portal_application_detail(application_id)
    except KeyError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Recruitment application not found") from exc


@router.delete("/applications/{application_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_recruitment_application_record(application_id: int, principal: Principal = Depends(require_permissions("recruitment:write"))) -> None:
    try:
        delete_recruitment_application(application_id)
    except KeyError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Recruitment application not found") from exc


@router.post("/applications/import", response_model=RecruitApplicationImportResult)
async def import_recruitment_application_records(
    plan_id: int = Form(...),
    file: UploadFile = File(...),
    principal: Principal = Depends(require_permissions("recruitment:write")),
) -> RecruitApplicationImportResult:
    try:
        rows = parse_recruitment_template(await file.read())
        return import_recruitment_applications(plan_id=plan_id, rows=rows, principal=principal)
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc


@router.get("/camp-offers", response_model=CampOfferListResponse)
def list_camp_offer_records(
    keyword: str | None = Query(default=None),
    plan_id: int | None = Query(default=None),
    is_sent_mail: bool | None = Query(default=None),
    is_agree: bool | None = Query(default=None),
    first_choice_advisor: str | None = Query(default=None),
    first_choice_team: str | None = Query(default=None),
    first_choice_score_op: str | None = Query(default=None, pattern="^(eq|ne|gt|ge|lt|le)$"),
    first_choice_score: float | None = Query(default=None),
    second_choice_advisor: str | None = Query(default=None),
    second_choice_team: str | None = Query(default=None),
    second_choice_score_op: str | None = Query(default=None, pattern="^(eq|ne|gt|ge|lt|le)$"),
    second_choice_score: float | None = Query(default=None),
    sort_by: str | None = Query(default=None),
    sort_order: str | None = Query(default=None),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=10, ge=1, le=1000),
    principal: Principal = Depends(require_permissions("recruitment_camp_offer:read")),
) -> CampOfferListResponse:
    return get_camp_offer_list(
        keyword=keyword,
        plan_id=plan_id,
        is_sent_mail=is_sent_mail,
        is_agree=is_agree,
        first_choice_advisor=first_choice_advisor,
        first_choice_team=first_choice_team,
        first_choice_score_op=first_choice_score_op,
        first_choice_score=first_choice_score,
        second_choice_advisor=second_choice_advisor,
        second_choice_team=second_choice_team,
        second_choice_score_op=second_choice_score_op,
        second_choice_score=second_choice_score,
        sort_by=sort_by,
        sort_order=sort_order,
        page=page,
        page_size=page_size,
        principal=principal,
    )


# --------------------------------------------------------------------------- #
# Offer-mail template management (system builtin + user uploaded .md files)
# --------------------------------------------------------------------------- #


def _builtin_offer_template_record(key: str) -> OfferTemplateRecord:
    path = resolve_builtin_offer_template_path(key)
    size = path.stat().st_size if path.exists() else 0
    return OfferTemplateRecord(
        id=key,
        filename=BUILTIN_OFFER_TEMPLATE_FILENAME[key],
        display_name=BUILTIN_OFFER_TEMPLATE_LABEL[key],
        size_bytes=size,
        uploaded_at=None,
        uploaded_by=None,
        is_builtin=True,
        source="builtin",
        builtin_key=key,
    )


def _uploaded_offer_template_record(path: Path) -> OfferTemplateRecord:
    stat = path.stat()
    # Filename convention: offer-{uuid-or-int}.md. We surface the raw
    # filename (basename) so reviewers can find the original upload.
    template_id = path.stem.removeprefix("offer-") or path.stem
    return OfferTemplateRecord(
        id=template_id,
        filename=path.name,
        display_name=path.name,
        size_bytes=int(stat.st_size),
        uploaded_at=datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
        uploaded_by=None,
        is_builtin=False,
        source="uploaded",
        builtin_key=None,
    )


def _coerce_template_id(raw: str) -> str:
    text = str(raw or "").strip()
    if not text:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="模板 id 不能为空",
        )
    if text.lower() in {"first", "second"}:
        return text.lower()
    if all(ch.isalnum() or ch in {"-", "_"} for ch in text):
        return text
    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail="模板 id 必须是 first / second 或字母数字组合",
    )


def _resolve_template_file(template_id: str) -> Path:
    if template_id in {"first", "second"}:
        return resolve_builtin_offer_template_path(template_id)
    return OFFER_TEMPLATE_UPLOAD_DIR / f"offer-{template_id}.md"


@router.get("/camp-offers/stats", response_model=CampOfferStats)
def get_camp_offer_stats_endpoint(
    keyword: str | None = Query(default=None),
    plan_id: int | None = Query(default=None),
    is_sent_mail: bool | None = Query(default=None),
    is_agree: bool | None = Query(default=None),
    first_choice_advisor: str | None = Query(default=None),
    first_choice_team: str | None = Query(default=None),
    first_choice_score_op: str | None = Query(default=None, pattern="^(eq|ne|gt|ge|lt|le)$"),
    first_choice_score: float | None = Query(default=None),
    second_choice_advisor: str | None = Query(default=None),
    second_choice_team: str | None = Query(default=None),
    second_choice_score_op: str | None = Query(default=None, pattern="^(eq|ne|gt|ge|lt|le)$"),
    second_choice_score: float | None = Query(default=None),
    principal: Principal = Depends(require_permissions("recruitment_camp_offer:read")),
) -> CampOfferStats:
    return get_camp_offer_stats(
        keyword=keyword,
        plan_id=plan_id,
        is_sent_mail=is_sent_mail,
        is_agree=is_agree,
        first_choice_advisor=first_choice_advisor,
        first_choice_team=first_choice_team,
        first_choice_score_op=first_choice_score_op,
        first_choice_score=first_choice_score,
        second_choice_advisor=second_choice_advisor,
        second_choice_team=second_choice_team,
        second_choice_score_op=second_choice_score_op,
        second_choice_score=second_choice_score,
    )


@router.get("/camp-offers/templates", response_model=OfferTemplateListResponse)
def list_camp_offer_templates(
    principal: Principal = Depends(require_permissions("recruitment_camp_offer:read")),
) -> OfferTemplateListResponse:
    items: list[OfferTemplateRecord] = [
        _builtin_offer_template_record("first"),
        _builtin_offer_template_record("second"),
    ]
    if OFFER_TEMPLATE_UPLOAD_DIR.exists():
        for child in sorted(OFFER_TEMPLATE_UPLOAD_DIR.iterdir(), key=lambda p: p.name):
            if not child.is_file():
                continue
            if child.suffix.lower() not in OFFER_TEMPLATE_ALLOWED_SUFFIXES:
                continue
            if not child.stem.startswith("offer-"):
                continue
            items.append(_uploaded_offer_template_record(child))
    return OfferTemplateListResponse(items=items)


@router.post("/camp-offers/templates", response_model=OfferTemplateRecord)
async def upload_camp_offer_template(
    file: UploadFile = File(...),
    principal: Principal = Depends(require_permissions("recruitment_camp_offer:write")),
) -> OfferTemplateRecord:
    original_name = Path(file.filename or "template.md").name
    suffix = Path(original_name).suffix.lower()
    if suffix not in OFFER_TEMPLATE_ALLOWED_SUFFIXES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="仅支持上传 .md / .markdown 文件",
        )
    content = await file.read(OFFER_TEMPLATE_MAX_BYTES + 1)
    if len(content) > OFFER_TEMPLATE_MAX_BYTES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="模板文件大小不能超过 1 MB",
        )
    try:
        text = content.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="模板文件必须为 UTF-8 编码",
        ) from exc
    if not text.strip():
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="模板文件不能为空",
        )

    OFFER_TEMPLATE_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    target_name = f"offer-{uuid4().hex}{suffix}"
    target_path = OFFER_TEMPLATE_UPLOAD_DIR / target_name
    target_path.write_text(text, encoding="utf-8")
    return _uploaded_offer_template_record(target_path)


@router.get("/camp-offers/templates/{template_id}/content")
def get_camp_offer_template_content(
    template_id: str = PathParam(...),
    principal: Principal = Depends(require_permissions("recruitment_camp_offer:read")),
) -> Response:
    resolved_id = _coerce_template_id(template_id)
    path = _resolve_template_file(resolved_id)
    if not path.exists():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"找不到模板文件: {path}",
        )
    text = path.read_text(encoding="utf-8")
    return Response(content=text, media_type="text/markdown; charset=utf-8")


@router.get("/camp-offers/templates/{template_id}/preview")
def get_camp_offer_template_preview(
    template_id: str = PathParam(...),
    principal: Principal = Depends(require_permissions("recruitment_camp_offer:read")),
) -> Response:
    resolved_id = _coerce_template_id(template_id)
    path = _resolve_template_file(resolved_id)
    if not path.exists():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"找不到模板文件: {path}",
        )
    text = path.read_text(encoding="utf-8")
    html = render_with_sample_placeholders(text)
    return Response(content=html, media_type="text/html; charset=utf-8")


@router.delete("/camp-offers/templates/{template_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_camp_offer_template(
    template_id: str = PathParam(...),
    principal: Principal = Depends(require_permissions("recruitment_camp_offer:write")),
) -> None:
    resolved_id = _coerce_template_id(template_id)
    if resolved_id in {"first", "second"}:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="内置模板不允许删除",
        )
    path = _resolve_template_file(resolved_id)
    if not path.exists():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"找不到模板文件: {path}",
        )
    try:
        path.unlink()
    except OSError as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"删除模板失败: {exc}",
        ) from exc


# Excel columns mirroring the CampOfferRecord surface, in the order
# operators see on screen.
_CAMP_OFFER_EXPORT_COLUMNS: list[tuple[str, str]] = [
    ("candidate_no", "报名号"),
    ("student_name", "姓名"),
    ("plan_name", "招生计划"),
    ("student_email", "邮箱"),
    ("student_phone", "手机号"),
    ("first_choice_advisor_name", "第一志愿导师"),
    ("first_choice_advisor_team_name", "第一志愿团队"),
    ("first_choice_screening_score", "第一志愿评分"),
    ("second_choice_advisor_name", "第二志愿导师"),
    ("second_choice_advisor_team_name", "第二志愿团队"),
    ("second_choice_screening_score", "第二志愿评分"),
    ("is_agree", "是否同意入营"),
    ("is_sent_mail", "是否已发邮件"),
    ("reason", "原因"),
    ("student_offer_submitted_at", "学生确认时间"),
    ("created_at", "创建时间"),
]


def _build_camp_offer_export_workbook(rows: list[dict[str, Any]]) -> bytes:
    from openpyxl import Workbook  # local import: keeps cold-start light

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "入营名单"
    sheet.append([label for _, label in _CAMP_OFFER_EXPORT_COLUMNS])
    for row in rows:
        sheet.append([row.get(key) for key, _ in _CAMP_OFFER_EXPORT_COLUMNS])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@router.get("/camp-offers/export")
def export_camp_offer_records(
    keyword: str | None = Query(default=None),
    plan_id: int | None = Query(default=None),
    is_sent_mail: bool | None = Query(default=None),
    is_agree: bool | None = Query(default=None),
    first_choice_advisor: str | None = Query(default=None),
    first_choice_team: str | None = Query(default=None),
    first_choice_score_op: str | None = Query(default=None, pattern="^(eq|ne|gt|ge|lt|le)$"),
    first_choice_score: float | None = Query(default=None),
    second_choice_advisor: str | None = Query(default=None),
    second_choice_team: str | None = Query(default=None),
    second_choice_score_op: str | None = Query(default=None, pattern="^(eq|ne|gt|ge|lt|le)$"),
    second_choice_score: float | None = Query(default=None),
    principal: Principal = Depends(require_permissions("recruitment_camp_offer:read")),
) -> StreamingResponse:
    rows = export_camp_offers(
        keyword=keyword,
        plan_id=plan_id,
        is_sent_mail=is_sent_mail,
        is_agree=is_agree,
        first_choice_advisor=first_choice_advisor,
        first_choice_team=first_choice_team,
        first_choice_score_op=first_choice_score_op,
        first_choice_score=first_choice_score,
        second_choice_advisor=second_choice_advisor,
        second_choice_team=second_choice_team,
        second_choice_score_op=second_choice_score_op,
        second_choice_score=second_choice_score,
        principal=principal,
    )
    content = _build_camp_offer_export_workbook(rows)
    filename = "入营名单_" + datetime.now().strftime("%Y%m%d%H%M%S") + ".xlsx"
    encoded_filename = quote(filename)
    return StreamingResponse(
        iter([content]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )


@router.get("/camp-offers/{offer_id}", response_model=CampOfferRecord)
def get_camp_offer_record_detail(
    offer_id: int,
    principal: Principal = Depends(require_permissions("recruitment_camp_offer:read")),
) -> CampOfferRecord:
    try:
        return get_camp_offer_detail(offer_id, principal=principal)
    except KeyError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Camp offer not found") from exc


@router.post("/camp-offers", response_model=CampOfferRecord, status_code=status.HTTP_201_CREATED)
def create_camp_offer_record(
    payload: CampOfferUpsert,
    principal: Principal = Depends(require_permissions("recruitment_camp_offer:write")),
) -> CampOfferRecord:
    try:
        return create_camp_offer(payload, principal=principal)
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc


@router.put("/camp-offers/{offer_id}", response_model=CampOfferRecord)
def update_camp_offer_record(
    offer_id: int,
    payload: CampOfferUpsert,
    principal: Principal = Depends(require_permissions("recruitment_camp_offer:write")),
) -> CampOfferRecord:
    try:
        return update_camp_offer(offer_id, payload, principal=principal)
    except KeyError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Camp offer not found") from exc
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc


@router.delete("/camp-offers/{offer_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_camp_offer_record(
    offer_id: int,
    principal: Principal = Depends(require_permissions("recruitment_camp_offer:write")),
) -> None:
    try:
        delete_camp_offer(offer_id, principal=principal)
    except KeyError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Camp offer not found") from exc

# ------------------------------------------------------------------
# 2026-07-03: 黑客松入取状态变更端点 (3 个独立端点，共用 service)
# - accept  -> accepted = "accepted_pending_send"  (录取未发送)
# - decline -> accepted = "declined"             (未录取)
# - pending -> accepted = "pending"              (待定)
# 权限: 需 recruitment_camp_offer:read (因为 service 层会基于导师/中心负责人身份做判断)。
# 状态可逆: 允许反复修改。
# ------------------------------------------------------------------
@router.post(
    "/camp-offers/{offer_id}/accept",
    response_model=CampOfferRecord,
)
def accept_camp_offer(
    offer_id: int,
    principal: Principal = Depends(require_permissions("recruitment_camp_offer:read")),
) -> CampOfferRecord:
    """2026-07-03: 录取学生 (设置 accepted=accepted_pending_send)。

    业务侧: 导师/中心负责人/书院管理员 在该学生 一/二志愿分数 >= 80 分时可点击。
    实际权限校验由 service 层 (_principal_can_change_camp_offer_accepted) 执行。
    """
    try:
        return set_camp_offer_accepted_status(
            offer_id, "accepted_pending_send", principal=principal
        )
    except KeyError as exc:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Camp offer not found"
        ) from exc
    except PermissionError as exc:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN, detail=str(exc)
        ) from exc
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)
        ) from exc


@router.post(
    "/camp-offers/{offer_id}/decline",
    response_model=CampOfferRecord,
)
def decline_camp_offer(
    offer_id: int,
    principal: Principal = Depends(require_permissions("recruitment_camp_offer:read")),
) -> CampOfferRecord:
    """2026-07-03: 不录取学生 (设置 accepted=declined)。

    业务侧: 导师/中心负责人/书院管理员 在该学生 一/二志愿分数 >= 80 分时可点击。
    实际权限校验由 service 层 (_principal_can_change_camp_offer_accepted) 执行。
    """
    try:
        return set_camp_offer_accepted_status(
            offer_id, "declined", principal=principal
        )
    except KeyError as exc:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Camp offer not found"
        ) from exc
    except PermissionError as exc:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN, detail=str(exc)
        ) from exc
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)
        ) from exc


@router.post(
    "/camp-offers/{offer_id}/pending",
    response_model=CampOfferRecord,
)
def mark_camp_offer_pending(
    offer_id: int,
    principal: Principal = Depends(require_permissions("recruitment_camp_offer:read")),
) -> CampOfferRecord:
    """2026-07-03: 待定学生 (设置 accepted=pending)。

    业务侧: 导师/中心负责人/书院管理员 在该学生 一/二志愿分数 >= 80 分时可点击。
    实际权限校验由 service 层 (_principal_can_change_camp_offer_accepted) 执行。
    """
    try:
        return set_camp_offer_accepted_status(
            offer_id, "pending", principal=principal
        )
    except KeyError as exc:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Camp offer not found"
        ) from exc
    except PermissionError as exc:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN, detail=str(exc)
        ) from exc
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)
        ) from exc



@router.post("/camp-offers/import", response_model=CampOfferImportResult)
async def import_camp_offer_records(
    plan_id: int | None = Form(default=None),
    file: UploadFile = File(...),
    principal: Principal = Depends(require_permissions("recruitment_camp_offer:write")),
) -> CampOfferImportResult:
    try:
        return import_camp_offers_from_excel(await file.read(), plan_id=plan_id)
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc


# 2026-07-03: 黑客松夏令营「评分导入」专用端点
# 区别于 /camp-offers/import (用 candidate_no 匹配入营名单), 本端点:
#   - 通过 dtlms_recruitment_applications.student_phone + student_email 联合匹配 (Q1+Q2)
#   - 仅更新 hackathon_score / hackathon_comments 两个字段 (Q4: 无条件覆盖)
#   - 匹配不到入营名单的行跳过, 在 issues 中报告 (Q3)
@router.post("/camp-offers/import-hackathon-scores", response_model=HackathonScoreImportResult)
async def import_hackathon_score_records(
    file: UploadFile = File(...),
    principal: Principal = Depends(require_permissions("recruitment_camp_offer:write")),
) -> HackathonScoreImportResult:
    try:
        return import_hackathon_scores_from_excel(await file.read())
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc

@router.post("/camp-offers/notify", response_model=CampOfferNotificationSendResponse)
def send_camp_offer_notification_records(
    payload: CampOfferNotificationSendRequest,
    principal: Principal = Depends(require_permissions("recruitment_camp_offer:write")),
) -> CampOfferNotificationSendResponse:
    try:
        return send_camp_offer_notifications(
            candidate_nos=payload.candidate_nos,
            choice=payload.choice,
            simulate=payload.simulate,
            simulate_recipient=payload.simulate_recipient,
            template_id=payload.template_id,
        )
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
    except RuntimeError as exc:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(exc)) from exc


@router.post("/camp-offers/confirm")
def confirm_camp_offer_record(
    email: str = Form(...),
    password: str = Form(...),
    choice: str = Form(...),
) -> dict[str, Any]:
    try:
        return submit_camp_offer_confirmation(email=email, password=password, choice=choice)
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc


@router.post("/advisor-screening/export-jobs", response_model=RegisteredPortalStudentExportJobCreateResponse)
def create_advisor_screening_export_job(
    payload: RegisteredPortalStudentExportRequest,
    principal: Principal = Depends(require_permissions("recruitment_advisor_screening:read")),
) -> RegisteredPortalStudentExportJobCreateResponse:
    try:
        scoped_payload = payload.model_copy(
            update={"ids": [], "export_scope": str(payload.export_scope or "advisor_screening").strip()}
        )
        return create_registered_portal_student_export_job(scoped_payload, principal=principal)
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
